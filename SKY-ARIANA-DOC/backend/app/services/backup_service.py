import os
import shutil
import json
from datetime import datetime
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session
from app.core.config import settings
from app.models.models import Invoice, Customer, Product, InvoiceItem, ShipmentDetails
from decimal import Decimal

BACKUP_DIR = os.path.join(settings.UPLOAD_DIR, "backups")

def create_db_backup() -> str:
    os.makedirs(BACKUP_DIR, exist_ok=True)
    db_file = settings.DATABASE_URL.replace("sqlite:///", "")
    if not os.path.exists(db_file):
        raise FileNotFoundError("Active database file not found")
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_filename = f"invoices_backup_{timestamp}.db"
    backup_path = os.path.join(BACKUP_DIR, backup_filename)
    
    shutil.copy2(db_file, backup_path)
    return backup_path

def restore_db_backup(backup_path: str) -> None:
    db_file = settings.DATABASE_URL.replace("sqlite:///", "")
    if not os.path.exists(backup_path):
        raise FileNotFoundError("Backup file not found")
    
    # Overwrite database file
    shutil.copy2(backup_path, db_file)

def export_invoices_to_json(db: Session) -> str:
    invoices = db.query(Invoice).all()
    data = []
    for inv in invoices:
        items = []
        for it in inv.items:
            items.append({
                "description": it.description,
                "hs_code": it.hs_code,
                "quantity": it.quantity,
                "unit": it.unit,
                "unit_price": float(it.unit_price),
                "amount": float(it.amount),
                "sort_order": it.sort_order
            })
            
        ship = {}
        if inv.shipment_details:
            s = inv.shipment_details
            ship = {
                "booking_number": s.booking_number,
                "bill_of_lading_number": s.bill_of_lading_number,
                "container_number": s.container_number,
                "container_type": s.container_type,
                "seal_number": s.seal_number,
                "commodity": s.commodity,
                "gross_weight": float(s.gross_weight) if s.gross_weight else None,
                "net_weight": float(s.net_weight) if s.net_weight else None,
                "package_count": s.package_count,
                "package_type": s.package_type,
                "port_of_loading": s.port_of_loading,
                "port_of_discharge": s.port_of_discharge,
                "vessel_name": s.vessel_name,
                "voyage_number": s.voyage_number,
                "shipping_line": s.shipping_line,
                "incoterms": s.incoterms,
                "etd": s.etd.isoformat() if s.etd else None,
                "eta": s.eta.isoformat() if s.eta else None
            }
            
        data.append({
            "invoice_number": inv.invoice_number,
            "invoice_date": inv.invoice_date.isoformat(),
            "due_date": inv.due_date.isoformat(),
            "customer_name": inv.customer.company_name,
            "currency": inv.currency,
            "status": inv.status,
            "invoice_template": inv.invoice_template or "premium_afghan_glass",
            "subtotal": float(inv.subtotal),
            "freight": float(inv.freight),
            "insurance": float(inv.insurance),
            "other_charges": float(inv.other_charges),
            "discount": float(inv.discount),
            "tax": float(inv.tax),
            "documentation_fees": float(inv.documentation_fees or 0.00),
            "customs_clearance_fees": float(inv.customs_clearance_fees or 0.00),
            "grand_total": float(inv.grand_total),
            "paid_amount": float(inv.paid_amount),
            "remaining_balance": float(inv.remaining_balance),
            "notes": inv.notes,
            "notify_party": inv.notify_party,
            "shipper_exporter": inv.shipper_exporter,
            "consignee_buyer": inv.consignee_buyer,
            "notify_party_enabled": inv.notify_party_enabled,
            "notify_party_same_as_consignee": inv.notify_party_same_as_consignee,
            "notify_party_customer_id": inv.notify_party_customer_id,
            "items": items,
            "shipment_details": ship
        })
        
    export_path = os.path.join(settings.UPLOAD_DIR, "exports", "invoices_export.json")
    os.makedirs(os.path.dirname(export_path), exist_ok=True)
    with open(export_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=4, ensure_ascii=False)
        
    return export_path

def export_invoices_to_excel(db: Session) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"
    
    headers = [
        "Invoice Number", "Invoice Date", "Due Date", "Customer", "Currency", "Status",
        "Template", "Subtotal", "Freight", "Insurance", "Other Charges", "Discount", "Tax",
        "Documentation Fees", "Customs Clearance Fees",
        "Grand Total", "Paid Amount", "Remaining Balance",
        "Notify Party Same As Consignee", "Notify Party JSON"
    ]
    ws.append(headers)
    
    invoices = db.query(Invoice).all()
    for inv in invoices:
        ws.append([
            inv.invoice_number,
            inv.invoice_date.isoformat(),
            inv.due_date.isoformat(),
            inv.customer.company_name,
            inv.currency,
            inv.status,
            inv.invoice_template or "premium_afghan_glass",
            float(inv.subtotal),
            float(inv.freight),
            float(inv.insurance),
            float(inv.other_charges),
            float(inv.discount),
            float(inv.tax),
            float(inv.documentation_fees or 0.00),
            float(inv.customs_clearance_fees or 0.00),
            float(inv.grand_total),
            float(inv.paid_amount),
            float(inv.remaining_balance),
            int(inv.notify_party_same_as_consignee or 0),
            json.dumps(inv.notify_party) if inv.notify_party else ""
        ])
        
    export_path = os.path.join(settings.UPLOAD_DIR, "exports", "invoices_export.xlsx")
    os.makedirs(os.path.dirname(export_path), exist_ok=True)
    wb.save(export_path)
    return export_path

def import_customers_from_excel(db: Session, file_path: str) -> int:
    wb = load_workbook(file_path, read_only=True)
    ws = wb.active
    
    imported_count = 0
    first_row = True
    for row in ws.iter_rows(values_only=True):
        if first_row:
            first_row = False
            continue
        
        company_name, contact_person, address, country, phone, email, tax_number = row[:7]
        if not company_name:
            continue
            
        customer = Customer(
            company_name=str(company_name),
            contact_person=str(contact_person) if contact_person else None,
            address=str(address) if address else None,
            country=str(country) if country else None,
            phone=str(phone) if phone else None,
            email=str(email) if email else None,
            tax_number=str(tax_number) if tax_number else None
        )
        db.add(customer)
        imported_count += 1
        
    db.commit()
    return imported_count

def import_products_from_excel(db: Session, file_path: str) -> int:
    wb = load_workbook(file_path, read_only=True)
    ws = wb.active
    
    imported_count = 0
    first_row = True
    for row in ws.iter_rows(values_only=True):
        if first_row:
            first_row = False
            continue
            
        name, description, hs_code, default_unit, default_price = row[:5]
        if not name:
            continue
            
        product = Product(
            name=str(name),
            description=str(description) if description else None,
            hs_code=str(hs_code) if hs_code else None,
            default_unit=str(default_unit) if default_unit else "PCS",
            default_price=Decimal(str(default_price)) if default_price else Decimal("0.00")
        )
        db.add(product)
        imported_count += 1
        
    db.commit()
    return imported_count
